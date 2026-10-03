"""Reading a duty roster from a spreadsheet, and writing it in one go.

The syndicate and the health directorate publish the night roster as a table. Typing it in
shift by shift is where mistakes come from, so an operator uploads the table instead. It is
read in full and checked row by row before anything is written; the result is a preview with
every problem named by its row, and applying it writes all of it or none of it
(`upsert_duty_shifts`). Re-applying the same file changes nothing.

Columns, in Arabic or English, in any order:

* the pharmacy — `facilityId`, or `pharmacy`/`الصيدلية` (its name in the province, spelling
  folded as search folds it), or `phone`/`الهاتف`;
* when — `date`/`التاريخ` with `from`/`من` and `to`/`إلى` in Damascus time (an end at or before
  the start is the next morning), or `startsAt` and `endsAt` as ISO date-times.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from typing import IO, Any
from uuid import UUID

from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime, parse_time

from business_hours.services import DAMASCUS
from facilities.models import Facility
from search.arabic import normalize_arabic

MAX_ROWS = 2000

# Every spelling a column may arrive under, folded to one key.
COLUMNS = {
    "facilityid": "facility_id",
    "facility": "facility_id",
    "pharmacy": "name",
    "name": "name",
    "الصيدلية": "name",
    "الصيدليه": "name",
    "اسم الصيدلية": "name",
    "phone": "phone",
    "الهاتف": "phone",
    "date": "date",
    "التاريخ": "date",
    "from": "from",
    "من": "from",
    "to": "to",
    "الى": "to",
    "إلى": "to",
    "startsat": "starts_at",
    "endsat": "ends_at",
}


@dataclass
class RowResult:
    line: int
    facility_id: str | None = None
    facility_name: str | None = None
    starts_at: datetime | None = None
    ends_at: datetime | None = None
    errors: list[str] = field(default_factory=list)


_FOLDED = {
    (key if key.isascii() else normalize_arabic(key)): value for key, value in COLUMNS.items()
}


def _header(value: Any) -> str | None:
    text = str(value or "").strip()
    if text.isascii():
        return _FOLDED.get(text.lower().replace("_", "").replace(" ", ""))
    return _FOLDED.get(normalize_arabic(text))


def read_table(upload: IO[bytes], filename: str) -> list[dict[str, Any]]:
    """The data rows of a CSV or XLSX file, keyed by canonical column."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        rows = _xlsx_rows(upload)
    else:
        rows = _csv_rows(upload)
    iterator = iter(rows)
    header = next(iterator, None)
    if header is None:
        raise ValueError("الملف فارغ.")
    keys = [_header(cell) for cell in header]
    if not any(keys):
        raise ValueError(
            "لم أتعرّف على أعمدة الملف. استعمل: الصيدلية، التاريخ، من، إلى (أو facilityId, "
            "date, from, to)."
        )
    table: list[dict[str, Any]] = []
    for values in iterator:
        if all(value in (None, "") for value in values):
            continue
        table.append({key: value for key, value in zip(keys, values, strict=False) if key})
        if len(table) > MAX_ROWS:
            raise ValueError(f"الملف أكبر من {MAX_ROWS} صف. قسّمه إلى ملفات أصغر.")
    return table


def _csv_rows(upload: IO[bytes]) -> Iterator[list[Any]]:
    text = io.TextIOWrapper(upload, encoding="utf-8-sig", newline="")
    sample = text.read(4096)
    text.seek(0)
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    yield from csv.reader(text, dialect)


def _xlsx_rows(upload: IO[bytes]) -> Iterator[list[Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(upload, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    for row in sheet.iter_rows(values_only=True):
        yield list(row)


def _safely(parse: Any, text: str) -> Any:
    """Django's parsers return None for a malformed value but raise for an impossible one
    (2026-13-40); a spreadsheet has both, and both are a row error, not a server error."""
    try:
        return parse(text)
    except ValueError:
        return None


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip().replace("/", "-")
    parsed: date | None = _safely(parse_date, text)
    if parsed is None and len(text.split("-")) == 3:
        day, month, year = text.split("-")
        if len(year) == 4:
            parsed = _safely(parse_date, f"{year}-{month.zfill(2)}-{day.zfill(2)}")
    return parsed


def _as_time(value: Any) -> time | None:
    if isinstance(value, time):
        return value
    if isinstance(value, datetime):
        return value.time()
    parsed: time | None = _safely(parse_time, str(value or "").strip())
    return parsed


def _as_moment(value: Any) -> datetime | None:
    moment: datetime | None
    if isinstance(value, datetime):
        moment = value
    else:
        moment = _safely(parse_datetime, str(value or "").strip()) if value else None
    if moment is None:
        return None
    return moment if timezone.is_aware(moment) else timezone.make_aware(moment, DAMASCUS)


class Directory:
    """The pharmacies a roster may name, looked up by id, folded name or phone."""

    def __init__(self, province_id: UUID) -> None:
        candidates = list(
            Facility.objects.filter(
                province_id=province_id,
                category__capabilities__supports_duty=True,
            ).exclude(status=Facility.Status.CLOSED)
        )
        self.by_id = {str(item.pk): item for item in candidates}
        self.by_name: dict[str, list[Facility]] = {}
        self.by_phone: dict[str, list[Facility]] = {}
        for item in candidates:
            self.by_name.setdefault(normalize_arabic(item.name_ar), []).append(item)
            if item.phone:
                self.by_phone.setdefault(_digits(item.phone), []).append(item)

    def find(self, row: dict[str, Any]) -> tuple[Facility | None, str | None]:
        if value := str(row.get("facility_id") or "").strip():
            found = self.by_id.get(value.lower())
            return (found, None) if found else (None, "لا صيدلية بهذا المعرّف في المحافظة.")
        if value := str(row.get("name") or "").strip():
            matches = self.by_name.get(normalize_arabic(value), [])
            if len(matches) == 1:
                return matches[0], None
            if not matches:
                return None, f"لا صيدلية باسم «{value}» في المحافظة."
            return None, f"أكثر من صيدلية باسم «{value}». أضف عمود الهاتف أو المعرّف."
        if value := _digits(str(row.get("phone") or "")):
            matches = self.by_phone.get(value, [])
            if len(matches) == 1:
                return matches[0], None
            return None, "لا صيدلية واحدة بهذا الهاتف في المحافظة."
        return None, "حدّد الصيدلية: الاسم أو الهاتف أو المعرّف."


def _digits(value: str) -> str:
    digits = "".join(ch for ch in value if ch.isdigit())
    return digits[-9:] if len(digits) >= 9 else digits


def check(table: Iterable[dict[str, Any]], province_id: UUID) -> list[RowResult]:
    """Every row read and checked against the province's pharmacies; nothing written."""
    directory = Directory(province_id)
    results: list[RowResult] = []
    for index, row in enumerate(table, start=2):
        result = RowResult(line=index)
        facility, problem = directory.find(row)
        if problem:
            result.errors.append(problem)
        elif facility is not None:
            result.facility_id = str(facility.pk)
            result.facility_name = facility.name_ar
        starts, ends = _window(row, result)
        result.starts_at, result.ends_at = starts, ends
        results.append(result)
    _flag_overlaps(results)
    return results


def _window(row: dict[str, Any], result: RowResult) -> tuple[datetime | None, datetime | None]:
    if row.get("starts_at") or row.get("ends_at"):
        starts, ends = _as_moment(row.get("starts_at")), _as_moment(row.get("ends_at"))
        if starts is None or ends is None:
            result.errors.append("بداية المناوبة أو نهايتها ليست تاريخاً ووقتاً صالحين.")
            return None, None
    else:
        day, opens, closes = (
            _as_date(row.get("date")),
            _as_time(row.get("from")),
            _as_time(row.get("to")),
        )
        if day is None:
            result.errors.append("التاريخ غير صالح؛ اكتبه هكذا: 2026-10-15.")
        if opens is None or closes is None:
            result.errors.append("وقت البداية أو النهاية غير صالح؛ اكتبه هكذا: 20:00.")
        if day is None or opens is None or closes is None:
            return None, None
        starts = timezone.make_aware(datetime.combine(day, opens), DAMASCUS)
        ends = timezone.make_aware(datetime.combine(day, closes), DAMASCUS)
        if ends <= starts:
            # A night shift: it ends the next morning.
            ends += timedelta(days=1)
    if ends <= starts:
        result.errors.append("النهاية قبل البداية.")
    elif ends - starts > timedelta(hours=48):
        result.errors.append("المناوبة أطول من ٤٨ ساعة؛ تحقق من التاريخين.")
    return starts, ends


def _flag_overlaps(results: list[RowResult]) -> None:
    """Two rows of the same pharmacy that overlap would be refused by the database anyway;
    naming them here says which lines to fix."""
    by_facility: dict[str, list[RowResult]] = {}
    for row in results:
        if row.facility_id and row.starts_at and row.ends_at and not row.errors:
            by_facility.setdefault(row.facility_id, []).append(row)
    for rows in by_facility.values():
        rows.sort(key=lambda item: item.starts_at or timezone.now())
        for earlier, later in zip(rows, rows[1:], strict=False):
            assert earlier.ends_at is not None and later.starts_at is not None
            if later.starts_at < earlier.ends_at and later.starts_at != earlier.starts_at:
                later.errors.append(f"تتداخل مع السطر {earlier.line} للصيدلية نفسها.")


class _Rollback(Exception):
    """Raised to undo a trial run once its outcome is known."""


def trial(results: list[RowResult]) -> dict[str, Any]:
    """What applying would do, row by row, against the shifts already stored; nothing kept.

    Each valid row is tried in its own savepoint, so one that clashes with a stored shift or
    a closure is named without hiding what the others would do.
    """
    from django.db import transaction

    from core.exceptions import DomainError

    from .services import upsert_duty_shifts

    outcomes: dict[int, str] = {}
    try:
        with transaction.atomic():
            for row in results:
                if row.errors:
                    continue
                try:
                    with transaction.atomic():
                        (done,) = upsert_duty_shifts([_row(row)], source="IMPORT")
                    outcomes[row.line] = done.outcome
                except DomainError as exc:
                    row.errors.append(exc.message)
            raise _Rollback
    except _Rollback:
        pass
    return summary(results, outcomes)


def _row(row: RowResult) -> dict[str, Any]:
    return {"facility_id": row.facility_id, "starts_at": row.starts_at, "ends_at": row.ends_at}


def apply(results: list[RowResult]) -> dict[str, Any]:
    """Write every row, all or nothing. The caller checked there are no errors."""
    from .services import upsert_duty_shifts

    done = upsert_duty_shifts([_row(row) for row in results], source="IMPORT")
    outcomes = {row.line: item.outcome for row, item in zip(results, done, strict=True)}
    return summary(results, outcomes)


def summary(results: list[RowResult], outcomes: dict[int, str]) -> dict[str, Any]:
    counts = {"CREATED": 0, "UPDATED": 0, "UNCHANGED": 0}
    for outcome in outcomes.values():
        counts[outcome] = counts.get(outcome, 0) + 1
    return {
        "rows": [
            {
                "line": row.line,
                "facilityId": row.facility_id,
                "facilityNameAr": row.facility_name,
                "startsAt": row.starts_at.isoformat() if row.starts_at else None,
                "endsAt": row.ends_at.isoformat() if row.ends_at else None,
                "outcome": outcomes.get(row.line),
                "problems": row.errors,
            }
            for row in results
        ],
        "errorCount": sum(1 for row in results if row.errors),
        "created": counts["CREATED"],
        "unchanged": counts["UNCHANGED"],
    }
